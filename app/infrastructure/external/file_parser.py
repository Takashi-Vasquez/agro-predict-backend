import csv
import io
from pathlib import Path
from uuid import UUID, uuid4

import openpyxl


def parse_csv(content: str) -> tuple[list[dict], list[str]]:
    reader = csv.DictReader(io.StringIO(content))
    if reader.fieldnames is None:
        return [], []
    headers = [h.strip() for h in reader.fieldnames]
    rows = []
    for i, row in enumerate(rows_from_csv(content), start=1):
        rows.append({"row_number": i, "data": {k.strip(): v.strip() if v else "" for k, v in row.items()}})
    return rows, headers


def rows_from_csv(content: str) -> list[dict]:
    reader = csv.DictReader(io.StringIO(content))
    return [row for row in reader]


def parse_xlsx(file_path: Path) -> tuple[list[dict], list[str]]:
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return [], []

    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None:
        return [], []

    headers = [str(h).strip() if h else "" for h in header_row]
    rows = []
    for i, row in enumerate(rows_iter, start=1):
        data = {}
        for j, h in enumerate(headers):
            val = row[j] if j < len(row) else None
            data[h] = str(val).strip() if val is not None else ""
        rows.append({"row_number": i, "data": data})

    wb.close()
    return rows, headers


def generate_csv_template(headers: list[str]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    return output.getvalue()


def generate_xlsx_template(file_path: Path, headers: list[str], field_labels: dict[str, str]) -> None:
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(file_path))
    ws = wb.add_worksheet("Datos")

    header_format = wb.add_format({
        "bold": True,
        "bg_color": "#D9E1F2",
        "border": 1,
        "text_wrap": True,
        "valign": "vcenter",
    })

    for col, h in enumerate(headers):
        label = field_labels.get(h, h)
        ws.write(0, col, label, header_format)
        ws.set_column(col, col, max(15, len(label) + 2))

    wb.close()
